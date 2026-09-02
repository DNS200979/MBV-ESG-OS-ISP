#!/usr/bin/env python3
"""
Carrega a lista oficial de ex-tarifários vigentes (MDIC) no catálogo.

Fonte: gov.br/mdic → SDIC → Ex-Tarifário → Estatísticas → "Ex-tarifários
vigentes" (planilha BIT e BK). Aceita o .xlsx original (requer openpyxl)
ou um CSV exportado dele:

    python3 tools/carregar_ex_tarifario.py "Ex-Tarifários Vigentes.xlsx" --dry-run
    python3 tools/carregar_ex_tarifario.py "Ex-Tarifários Vigentes.xlsx"

Os nomes de coluna padrão são os da planilha oficial (NCM, EX, DESCRIÇÃO,
Ato Legal, Início/Fim da vigência). Para CSV só é preciso a stdlib. Envia em lotes via Management API do Supabase
(token em $SUPABASE_ACCESS_TOKEN ou ~/.supabase-token). A escrita usa
ON CONFLICT (ncm, descricao) DO NOTHING — recarga é idempotente.

Tudo entra como validacao='nao_validado' (§12): a conferência do Ex
específico contra o bem importado continua sendo do despachante.
"""
import argparse, csv, json, os, re, subprocess, sys
from pathlib import Path

REF_PADRAO = "kyrivjhglgtxwecovtcd"
# NCMs de interesse do ecossistema telecom (prefixos); ajuste com --prefixos
PREFIXOS_PADRAO = "8517,8544,8471,8525,9001"


def token():
    t = os.environ.get("SUPABASE_ACCESS_TOKEN")
    if not t:
        p = Path.home() / ".supabase-token"
        if p.exists():
            t = p.read_text().strip()
    if not t:
        sys.exit("Defina SUPABASE_ACCESS_TOKEN ou grave o token em ~/.supabase-token")
    return t


def so_digitos(v):
    return re.sub(r"\D", "", v or "")


def ler_registros(caminho, aba=None):
    """Itera dicts coluna→valor a partir de .xlsx (openpyxl) ou .csv (stdlib)."""
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Para ler .xlsx instale openpyxl (pip install openpyxl) "
                     "ou exporte a planilha como CSV.")
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        ws = wb[aba] if aba else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(it)]
        for row in it:
            yield dict(zip(cab, row))
    else:
        with open(caminho, newline="", encoding="utf-8-sig") as f:
            yield from csv.DictReader(f)


def normaliza_data(v):
    v = (v or "").strip()
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", v)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.match(r"\d{4}-\d{2}-\d{2}", v)
    return m.group(0) if m else None


def executa_sql(ref, tok, sql):
    """Management API via curl — urllib é barrado pelo Cloudflare (erro 1010)."""
    corpo = json.dumps({"query": sql})
    r = subprocess.run(
        ["curl", "-s", "-o", "/dev/stderr", "-w", "%{http_code}", "-X", "POST",
         "-H", f"Authorization: Bearer {tok}",
         "-H", "Content-Type: application/json",
         "--data-binary", "@-",
         f"https://api.supabase.com/v1/projects/{ref}/database/query"],
        input=corpo.encode(), capture_output=True)
    codigo = r.stdout.decode().strip()
    if codigo not in ("200", "201"):
        sys.exit(f"Falha na API (HTTP {codigo}): {r.stderr.decode()[:300]}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("arquivo", help="planilha oficial do MDIC (.xlsx) ou CSV exportado dela")
    ap.add_argument("--col-ncm", default="NCM")
    ap.add_argument("--col-ex", default="EX", help="coluna do código do Ex (opcional)")
    ap.add_argument("--col-descricao", default="DESCRIÇÃO")
    ap.add_argument("--col-ato", default="Ato Legal", help="coluna da resolução/ato")
    ap.add_argument("--col-ini", default="Início da vigência", help="coluna da vigência inicial")
    ap.add_argument("--col-fim", default="Fim da vigência", help="coluna da vigência final")
    ap.add_argument("--aba", default=None, help="aba do xlsx (padrão: a primeira)")
    ap.add_argument("--prefixos", default=PREFIXOS_PADRAO,
                    help=f"prefixos de NCM a importar (padrão: {PREFIXOS_PADRAO}); use '' para todos")
    ap.add_argument("--ref", default=os.environ.get("SUPABASE_PROJECT_REF", REF_PADRAO))
    ap.add_argument("--lote", type=int, default=200)
    ap.add_argument("--dry-run", action="store_true", help="só mostra o que seria carregado")
    a = ap.parse_args()

    prefixos = tuple(p.strip() for p in a.prefixos.split(",") if p.strip())
    linhas, puladas = [], 0

    for reg in ler_registros(a.arquivo, a.aba):
        ncm = so_digitos(str(reg.get(a.col_ncm) or ""))
        desc = str(reg.get(a.col_descricao) or "").strip()
        if len(ncm) != 8 or not desc:
            puladas += 1
            continue
        if prefixos and not ncm.startswith(prefixos):
            puladas += 1
            continue
        ex = so_digitos(str(reg.get(a.col_ex) or "")) if a.col_ex else ""
        if ex:
            # o Ex é a identidade fina do pleito — vai no início da descrição
            desc = f"Ex {ex} — {desc}"
        linhas.append({
            "ncm": ncm,
            "descricao": desc[:500],
            "ato": (str(reg.get(a.col_ato) or "").strip()[:200] or None) if a.col_ato else None,
            "ini": normaliza_data(str(reg.get(a.col_ini) or "")) if a.col_ini else None,
            "fim": normaliza_data(str(reg.get(a.col_fim) or "")) if a.col_fim else None,
        })

    print(f"{len(linhas)} pleitos a carregar ({puladas} linhas fora do filtro).")
    if a.dry_run:
        for l in linhas[:10]:
            print("  ", l["ncm"], "—", l["descricao"][:70])
        return
    if not linhas:
        return

    tok = token()
    q = lambda v: "null" if v is None else "'" + v.replace("'", "''") + "'"
    for i in range(0, len(linhas), a.lote):
        vals = ",\n".join(
            f"({q(l['ncm'])}, {q(l['descricao'])}, {q(l['ato'])}, 0.00, {q(l['ini'])}, {q(l['fim'])}, 'nao_validado')"
            for l in linhas[i:i + a.lote])
        executa_sql(a.ref, tok, f"""
            insert into ex_tarifario_pleitos
              (ncm, descricao, ato_normativo, aliquota_ii_reduzida, vigencia_ini, vigencia_fim, validacao)
            values {vals}
            on conflict (ncm, descricao) do nothing;""")
        print(f"  lote {i // a.lote + 1}: ok")
    print("Concluído. Recarga é idempotente (ON CONFLICT DO NOTHING).")


if __name__ == "__main__":
    main()
